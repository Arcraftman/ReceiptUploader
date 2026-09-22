"""Reference management ratios; do not mislabel YTD data as certification tests."""
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from .rd_report import RD_TOTAL


def add_hightech_report(book, matches):
    s = book.create_sheet('高新企业相关指标监测', book.index(book['2026年利润和负债']))
    s.sheet_view.showGridLines = False
    s.freeze_panes = 'C6'
    for col,width in {'A':4,'B':8,'C':58,'D':20,'E':20,'F':25,'G':45,'H':4}.items():
        s.column_dimensions[col].width = width
    s.merge_cells('B2:G3')
    s['B2'] = '高新企业相关指标监测'
    s['B2'].font = Font(name='Arial',size=22,bold=True,color='FFFFFF')
    s['B2'].fill = PatternFill('solid',fgColor='17324D')
    s['B2'].alignment = Alignment(vertical='center')
    s.merge_cells('B4:G4')
    s['B4'] = '=IF(\'控制台\'!B5="","选择公司及月份并刷新",\'控制台\'!B4&"  |  截至 "&\'控制台\'!B5&"  |  原表当年累计管理监测口径")'
    headers = ['序号','指标','参考门槛','当年累计比值','参考比较','口径说明']
    for col,title in enumerate(headers,2):
        s.cell(6,col,title)
    s['B8'],s['B9'] = 2,3
    s['C8'] = '研究开发费用占销售收入比例\n正式条件：近三个会计年度合计研发费用÷同期销售收入；最近一年收入≤5,000万元为5%，≤2亿元为4%，超过2亿元为3%。'
    s['C9'] = '高新技术产品（服务）收入占总收入比例\n正式条件：近一年高新技术产品（服务）收入÷同期总收入，不低于60%。'
    s['B7'] = 1
    s['C7'] = '企业从事研发和相关技术创新活动的科技人员占企业当年职工总数的比例不低于10%'
    s['D7'] = 0.1
    s['D7'].number_format = '0.0%'
    income = "'2026年利润和负债'!AI5"
    tech = "'2026年利润和负债'!AI47"
    s['D8'] = f'=IF(AND({matches},ISNUMBER({income})),IF({income}<=50000000,5%,IF({income}<=200000000,4%,3%)),"")'
    s['D9'] = 0.6
    s['E8'] = f'=IF(AND({matches},ISNUMBER({RD_TOTAL}),ISNUMBER({income}),{income}>0),{RD_TOTAL}/{income},"")'
    s['E9'] = f'=IF(AND({matches},ISNUMBER({tech}),ISNUMBER({income}),{income}>0),{tech}/{income},"")'
    for row in (8,9):
        s[f'F{row}'] = f'=IF(OR(E{row}="",D{row}=""),"待补数据",IF(E{row}>=D{row},"本期参考比值达到门槛","本期参考比值低于门槛"))'
        for col in ('D','E'):
            s[f'{col}{row}'].number_format = '0.0%;[Red](0.0%);"—"'
    s['G8'] = '沿用原表：当年累计研发费用÷当年累计营业收入。门槛按本期累计收入暂作分档；不是近三年正式认定结果。历史年度与最近完整年度收入未补齐。'
    s['G9'] = '沿用原表：研发收入÷营业收入，仅作代理指标。研发收入不自动等同于全部高新收入；正式总收入与营业收入也可能不同。未映射收入则留空。'
    for row in range(6,10):
        s.row_dimensions[row].height = 32 if row==6 else 135
        for col in range(2,8):
            cell = s.cell(row,col)
            cell.font = Font(name='Arial',size=11,bold=row==6,color='FFFFFF' if row==6 else '243746')
            cell.fill = PatternFill('solid',fgColor='17324D' if row==6 else 'F0F5F9')
            cell.alignment = Alignment(wrap_text=True,vertical='center')
            cell.border = Border(bottom=Side(style='thin',color='D3DFE8'))
    s.merge_cells('B11:G12')
    s['B11'] = '比值随所选月份及已刷新数据自动计算。正式申报需补齐三年归集数据、确认高新收入分类及总收入口径。'
    s['B11'].alignment = Alignment(wrap_text=True,vertical='center')
    s['B11'].font = Font(name='Arial',size=11,color='536579')
    s.merge_cells('B14:G14')
    s['B14'] = '口径参考：国家税务总局《高新技术企业优惠情况及明细表》'
    s['B14'].hyperlink = 'https://12366.chinatax.gov.cn/bzds/039/039.html'
    s.print_area = 'A1:H15'
    s.page_setup.orientation = 'landscape'
    s.page_setup.paperSize = s.PAPERSIZE_A3
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.fitToWidth = s.page_setup.fitToHeight = 1
