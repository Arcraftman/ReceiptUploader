"""Formula-driven dashboard over the existing monthly profit snapshot."""
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, NumRef
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break

NAME = '财务分析看板'
FIRST = 83
MONEY = '#,##0.00;[Red](#,##0.00);"—"'


def add_dashboard(book, matches):
    sheet = book.create_sheet(NAME, 1)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 80
    sheet.freeze_panes = 'B83'
    sheet.sheet_properties.tabColor = '007F86'
    for col in range(1, 17):
        sheet.column_dimensions[get_column_letter(col)].width = 14.5
    for row in range(1, 101):
        sheet.row_dimensions[row].height = 23
        for col in range(1, 17):
            cell = sheet.cell(row, col)
            cell.font = Font(name='Arial', size=10, color='243746')
            cell.fill = PatternFill('solid', fgColor='F3F6FA')
    def band(address, text, size=11, color='17324D'):
        sheet.merge_cells(address)
        cell = sheet[address.split(':')[0]]
        cell.value = text
        cell.fill = PatternFill('solid', fgColor=color)
        cell.font = Font(name='Arial', size=size, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center')
    band('A1:P2', '财务分析看板  /  收入 · 成本 · 费用 · 利润', 22)
    sheet.merge_cells('A3:P3')
    sheet['A3'] = '=IF($R$1=0,"请在控制台选择公司和月份并刷新",\'控制台\'!B4&"  |  "&LEFT(\'控制台\'!B5,4)&"年1—"&$R$1&"月  |  金额：人民币元")'
    sheet['R1'] = f'=IFERROR(IF({matches},VALUE(RIGHT(\'控制台\'!B5,2)),0),0)'
    sheet.column_dimensions['R'].hidden = True
    # All KPIs require complete monthly observations. Missing source amounts stay blank.
    cards = [('A','D','累计收入','B',False), ('E','H','累计营业成本','C',False),
             ('I','L','累计期间费用','G',False), ('M','P','累计净利润','J',False)]
    for left, right, title, source, percent in cards:
        band(f'{left}5:{right}5', title)
        sheet.merge_cells(f'{left}6:{right}7')
        cell = sheet[f'{left}6']
        cell.value = f'=IF(AND($R$1>0,COUNT({source}83:{source}94)=$R$1),SUM({source}83:{source}94),"")'
        cell.number_format = MONEY
        cell.font = Font(name='Arial', size=23, bold=True, color='007F86')
    band('A9:D9', '累计毛利率')
    band('E9:H9', '累计净利率')
    sheet.merge_cells('A10:D10'); sheet.merge_cells('E10:H10')
    sheet['A10'] = '=IF(AND(ISNUMBER(A6),ISNUMBER(E6),A6<>0),(A6-E6)/A6,"")'
    sheet['E10'] = '=IF(AND(ISNUMBER(A6),ISNUMBER(M6),A6<>0),M6/A6,"")'
    for coord in ('A10', 'E10'):
        sheet[coord].number_format = '0.0%;[Red](0.0%);"—"'
        sheet[coord].font = Font(name='Arial', size=19, bold=True, color='007F86')
    sheet.merge_cells('I9:P10')
    sheet['I9'] = '=IF($R$1=0,"选择已变化或尚未刷新，暂不展示旧数据",IF(AND(ISNUMBER(A6),ISNUMBER(E6),ISNUMBER(I6),ISNUMBER(M6)),"期间数据完整 · 与下方明细联动","存在缺失项目：相关月份及累计指标留空，请核对利润表"))'
    sheet['I9'].alignment = Alignment(wrap_text=True, vertical='center')
    headers = ['月份','营业收入','营业成本','销售费用','管理费用','财务费用','期间费用',
               '毛利润','利润总额','净利润','毛利率','净利率','费用率','累计收入','累计净利润','收入环比']
    band('A80:P80', '月度明细  /  所有图表均来自下表；金额为人民币元')
    for i, header in enumerate(headers, 1):
        sheet.cell(82,i,header)
        sheet.cell(82,i).fill = PatternFill('solid', fgColor='17324D')
        sheet.cell(82,i).font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    raw = "'月度趋势'!"
    amounts, periods = raw+'$D$5:$D$1000', raw+'$A$5:$A$1000'
    sources = {'B':('E','营业收入'), 'C':('E','营业成本'), 'D':('C','*销售费用'),
               'E':('C','*管理费用'), 'F':('C','*财务费用'), 'I':('C','*利润总额*'), 'J':('E','净利润')}
    for row in range(83,95):
        month = row-82
        sheet[f'A{row}'] = f'=IF({month}<=$R$1,LEFT(\'控制台\'!B5,4)&"-{month:02}","")'
        for col, (source, label) in sources.items():
            criteria = f'{periods},$A{row},{raw}${source}$5:${source}$1000,"{label}"'
            sheet[f'{col}{row}'] = f'=IF(AND($A{row}<>"",COUNTIFS({criteria})=1,COUNTIFS({criteria},{amounts},"<>")=1),SUMIFS({amounts},{criteria}),"")'
        for col, expression, needed in [('G',f'SUM(D{row}:F{row})',f'COUNT(D{row}:F{row})=3'),
                                        ('H',f'B{row}-C{row}',f'COUNT(B{row}:C{row})=2')]:
            sheet[f'{col}{row}'] = f'=IF({needed},{expression},"")'
        for col, numerator in [('K','H'),('L','J'),('M','G')]:
            sheet[f'{col}{row}'] = f'=IF(AND(ISNUMBER({numerator}{row}),ISNUMBER(B{row}),B{row}<>0),{numerator}{row}/B{row},"")'
        for col, source in [('N','B'),('O','J')]:
            sheet[f'{col}{row}'] = f'=IF(AND(A{row}<>"",COUNT({source}$83:{source}{row})={month}),SUM({source}$83:{source}{row}),"")'
        sheet[f'P{row}'] = '=""' if month==1 else f'=IF(AND(ISNUMBER(B{row}),ISNUMBER(B{row-1}),B{row-1}>0),B{row}/B{row-1}-1,"")'
        for col in range(2,17):
            sheet.cell(row,col).number_format = '0.0%;[Red](0.0%);"—"' if col in (11,12,13,16) else MONEY
            sheet.cell(row,col).fill = PatternFill('solid', fgColor='FFFFFF' if row%2==0 else 'E7EFF5')
    # Excel expands these native chart references to exactly the selected month.
    # Formula-empty strings plot as zero in Excel charts. Hidden NA helpers on this
    # same sheet keep absent observations out of charts without another worksheet.
    for col in range(1,17):
        letter = get_column_letter(col)
        chart_letter = letter if col == 1 else get_column_letter(col+18)
        if col > 1:
            sheet.column_dimensions[chart_letter].hidden = True
            for row in range(83,95):
                sheet[f'{chart_letter}{row}'] = f'=IF(ISNUMBER({letter}{row}),{letter}{row},NA())'
        book.defined_names.add(DefinedName(f'FinanceDash_{letter}', localSheetId=1, attr_text=f"'{NAME}'!${chart_letter}$83:INDEX('{NAME}'!${chart_letter}$83:${chart_letter}$94,MAX(1,'{NAME}'!$R$1))"))
    specs = [('收入与营业成本',BarChart,[2,3],'A14',False),
             ('毛利润与净利润',LineChart,[8,10],'J14',False),
             ('期间费用构成',BarChart,[4,5,6],'A36',True),
             ('毛利率 · 净利率 · 费用率',LineChart,[11,12,13],'J36',False),
             ('累计收入与累计净利润',LineChart,[14,15],'A58',False),
             ('收入环比增长率',BarChart,[16],'J58',False)]
    for title, kind, columns, anchor, stacked in specs:
        chart = kind()
        is_ratio = columns[0] in (11,16)
        chart.title = title + ('（%）' if is_ratio else '（元）')
        chart.width, chart.height = 19.0, 10.6
        chart.style = 13
        chart.y_axis.numFmt = '0%' if is_ratio else '#,##0;[Red](#,##0);0'
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.tickLblPos = 'low'
        chart.y_axis.tickLblPos = 'nextTo'
        chart.title.overlay = False
        chart.x_axis.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1100)))])
        chart.y_axis.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1100)))])
        chart.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1500,b=True))
        chart.display_blanks = 'gap'
        chart.visible_cells_only = False
        if kind is BarChart:
            chart.type = 'col'
            chart.grouping = 'stacked' if stacked else 'clustered'
            if stacked: chart.overlap = 100
        if columns[0] in (11,16): chart.y_axis.numFmt = '0%'
        for i,col in enumerate(columns):
            chart.add_data(Reference(sheet,min_col=col,min_row=82,max_row=94),titles_from_data=True)
            series = chart.series[-1]
            series.val.numRef.f = f"'{NAME}'!FinanceDash_{get_column_letter(col)}"
            series.cat = AxDataSource(numRef=NumRef(f=f"'{NAME}'!FinanceDash_A"))
            color = ['007F86','E69B45','7284B0'][i%3]
            series.graphicalProperties.line.solidFill = color
            if kind is BarChart:
                series.graphicalProperties.solidFill = color
                series.invertIfNegative = False
        chart.legend.position = 'b'
        chart.legend.overlay = False
        chart.legend.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1100)))])
        if len(columns) == 1:
            chart.legend = None
            chart.varyColors = False
        sheet.add_chart(chart,anchor)
    notes = [
        '口径：期间费用＝销售费用＋管理费用＋财务费用；研发费用通常包含在管理费用中，不再重复加计。',
        '毛利润＝营业收入－营业成本；净利润、利润总额直接读取原利润表，保留税费及其他损益影响。',
        '累计比例按累计金额计算；缺项留空，零收入不算比例；上月收入≤0时不算环比；1月无环比。',
        '数据来源：本工作簿“月度趋势”（账无忧只读利润表快照）。切换公司或月份后请先刷新。',
        '布局参考：Microsoft 财务管理模板、Smartsheet 月度损益模板；参考仅用于布局，不作为财务数据。',
    ]
    explanations = [
        ('A29:G32', '统计口径：当月营业收入与营业成本。\n阅读重点：比较两者变化，观察收入增长是否伴随成本同步上升。两者差额为毛利润，尚未扣除期间费用及税费。'),
        ('J29:P32', '统计口径：毛利润＝收入－营业成本；净利润直接取利润表。\n阅读重点：比较毛利与净利的走势。差额包含期间费用、税费及其他损益，不等于期间费用。'),
        ('A51:G54', '统计口径：销售费用、管理费用、财务费用按月堆积。\n阅读重点：观察费用结构和异常月份。研发费用不重复加计；红字冲销或负费用保留在零线下方。'),
        ('J51:P54', '统计口径：毛利率、净利率、期间费用率均除以当月收入。\n阅读重点：比较盈利能力与费用负担。零收入或缺失时留空，累计卡片按累计金额计算。'),
        ('A73:G76', '统计口径：自1月起逐月累计收入和净利润，均为人民币元。\n阅读重点：观察累计规模和利润积累。净利润不是现金流；前期数据缺失时不显示不完整累计值。'),
        ('J73:P76', '统计口径：本月收入÷上月收入－1。\n阅读重点：正值表示较上月增长，负值表示下降。1月、上月收入≤0或任一月份缺失时不计算环比。'),
    ]
    for address, explanation in explanations:
        sheet.merge_cells(address)
        cell = sheet[address.split(':')[0]]
        cell.value = explanation
        cell.font = Font(name='Arial', size=12, color='40566B')
        cell.alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    # The XLSM installer overlays each caption area with a cell-linked text box.
    # Keep its text in formulas so changing company/month cannot show stale AI output.
    for col in ('AI','AJ','AK'):
        sheet.column_dimensions[col].hidden = True
    for index in range(1,7):
        row = index+1
        sheet[f'AI{row}'] = '尚未生成 DeepSeek 解读。刷新数据后自动生成，也可点击“生成图表解读”。'
        sheet[f'AJ{row}'] = f'=IF(AND($R$1>0,\'控制台\'!B4=$AJ$1,\'控制台\'!B5=$AK$1),AI{row},"公司或月份已变化，请刷新后生成解读。")'
    for address,_ in explanations:
        sheet[address.split(':')[0]] = None
    for row,note in enumerate(notes,96):
        sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=16)
        sheet.cell(row,1,note)
    sheet['A100'].hyperlink = 'https://excel.cloud.microsoft/create/en/financial-management-templates/'
    sheet.print_options.horizontalCentered = True
    sheet.print_area = 'A1:P100'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for row in (34,56,78):
        sheet.row_breaks.append(Break(id=row))
