"""Single visible R&D detail table with hidden source calculations."""
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SUBJECTS = {'差旅费':'43010101','交通费':'43010102','研发人员工资':'43010103',
            '研发人员公积金':'43010104','研发人员社保':'43010105','软件购置费':'43010106',
            '福利费':'43010107','折旧费':'43010108'}
RD_TOTAL = "'2026研发费用'!O15"


def add_rd_report(book, matches):
    s = book.create_sheet('2026研发费用', book.index(book['2026年利润和负债']))
    s.sheet_view.showGridLines = False
    s.sheet_view.zoomScale = 85
    s.freeze_panes = 'C7'
    s.column_dimensions['A'].width = 4
    s.column_dimensions['B'].width = 24
    for col in range(3,18):
        s.column_dimensions[get_column_letter(col)].width = 14
    s.column_dimensions['O'].width = 18
    s.column_dimensions['Q'].width = 4
    s.merge_cells('B2:P3')
    s['B2'] = '2026年研发费用明细'
    s['B2'].font = Font(name='Arial',size=22,bold=True,color='FFFFFF')
    s['B2'].fill = PatternFill('solid',fgColor='17324D')
    s['B2'].alignment = Alignment(vertical='center')
    s.merge_cells('B4:P4')
    s['B4'] = '=IF($AK$1=0,"请选择公司和月份并刷新",\'控制台\'!B4&"  |  2026年1—"&$AK$1&"月  |  人民币元  |  合计金额降序")'
    s['AK1'] = f'=IFERROR(IF(AND({matches},LEFT(\'控制台\'!B5,4)="2026"),VALUE(RIGHT(\'控制台\'!B5,2)),0),0)'
    raw = "'年度分析数据'!"
    a,b,c,d,e = [raw+f'${col}$5:${col}$20000' for col in 'ABCDE']
    for row,(label,code) in enumerate(SUBJECTS.items(),2):
        s[f'AR{row}'], s[f'AS{row}'] = label, code
        for month in range(1,13):
            col = get_column_letter(45+month)
            criteria = f'{a},"2026-{month:02}",{b},"subject_debit_leaf",{c},$AS{row},{d},"*"&$AR{row}'
            s[f'{col}{row}'] = f'=IF(AND($AK$1>={month},COUNTIFS({criteria})=1,COUNTIFS({criteria},{e},"<>")=1),SUMIFS({e},{criteria}),"")'
        s[f'BF{row}'] = f'=IF(COUNT(AT{row}:BE{row})>0,SUM(AT{row}:BE{row}),"")'
        s[f'BG{row}'] = f'=IF(ISNUMBER(BF{row}),RANK(BF{row},$BF$2:$BF$9,0)+COUNTIF($BF$2:BF{row},BF{row})-1,COUNT($BF$2:$BF$9)+COUNTBLANK($BF$2:BF{row}))'
    for col,title in enumerate(['研发费用明细',*[f'{m}月' for m in range(1,13)],'合计金额','占比'],2):
        s.cell(6,col,title)
    for row in range(7,15):
        s[f'B{row}'] = f'=INDEX($AR$2:$AR$9,MATCH(ROW()-6,$BG$2:$BG$9,0))'
        for month in range(1,13):
            col = get_column_letter(45+month)
            source = f'INDEX(${col}$2:${col}$9,MATCH($B{row},$AR$2:$AR$9,0))'
            s.cell(row,month+2,f'=IF(ISNUMBER({source}),{source},"")')
        s[f'O{row}'] = f'=IF(COUNT(C{row}:N{row})>0,SUM(C{row}:N{row}),"")'
        s[f'P{row}'] = f'=IF(AND(ISNUMBER(O{row}),ISNUMBER($O$15),$O$15<>0),O{row}/$O$15,"")'
    s['B15'] = '合计'
    for col in range(3,17):
        letter = get_column_letter(col)
        s[f'{letter}15'] = f'=IF(COUNT({letter}7:{letter}14)>0,SUM({letter}7:{letter}14),"")'
    for row in range(6,16):
        s.row_dimensions[row].height = 30
        for col in range(2,17):
            cell = s.cell(row,col)
            cell.font = Font(name='Arial',size=11,bold=row in (6,15),color='FFFFFF' if row==6 else '243746')
            cell.fill = PatternFill('solid',fgColor='17324D' if row==6 else 'DDEDF0' if row==15 else 'F0F5F9' if row%2 else 'FFFFFF')
            cell.border = Border(bottom=Side(style='thin',color='DAE3EB'))
            cell.alignment = Alignment(vertical='center',horizontal='left' if col==2 else 'right',shrinkToFit=True)
            if row>6 and col>2:
                cell.number_format = '0.0%;[Red](0.0%);"—"' if col==16 else '#,##0.00;[Red](#,##0.00);"—"'
    notes = ['取研发支出末级科目本期借方发生额，保留红字冲销，不扣减月末结转贷方。',
             '按年初至所选月份合计降序排列；后续月份保留模板并留空。占比以全部分类合计为分母。',
             '科目编码及名称同时匹配参考表映射；未匹配项目不计入。空白表示未读取到金额，并非自动填零。']
    for row,note in enumerate(notes,18):
        s.merge_cells(start_row=row,start_column=2,end_row=row,end_column=16)
        s.cell(row,2,note).font = Font(name='Arial',size=10,color='536579')
        s.row_dimensions[row].height = 25
    for col in range(37,60):
        s.column_dimensions[get_column_letter(col)].hidden = True
    s.print_area = 'A1:Q21'
    s.print_options.horizontalCentered = True
    s.page_setup.orientation = 'landscape'
    s.page_setup.paperSize = s.PAPERSIZE_A3
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.fitToWidth = s.page_setup.fitToHeight = 1
