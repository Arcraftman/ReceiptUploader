"""Build the blank finance workbook using only the installed Python package."""
from __future__ import annotations

import argparse
from copy import copy
from importlib.resources import files
import json
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from ..project_runtime import project_root
from .dashboard import add_dashboard
from .rd_report import add_rd_report
from .hightech_report import add_hightech_report

MONEY = '#,##0.00;[Red](#,##0.00);"—"'
MATCHES = "AND('控制台'!B4<>\"\",'控制台'!B5<>\"\",'控制台'!B4='刷新信息'!B6,'控制台'!B5='刷新信息'!B8)"
AS_OF = "DATE(VALUE(LEFT('控制台'!B5,4)),VALUE(RIGHT('控制台'!B5,2))+1,1)-1"
REPORT_YEAR = 2026


def add_annual_analysis_data(book: Workbook) -> None:
    """Create the hidden transport sheet used to fill the annual report."""
    sheet = book.create_sheet('年度分析数据')
    sheet.append(['年度分析数据'])
    sheet.append(['由只读刷新服务生成；不要人工编辑'])
    sheet.append([])
    sheet.append(['月份', '数据类型', '编码', '项目', '金额'])
    sheet.freeze_panes = 'A5'
    sheet.sheet_state = 'hidden'


def add_profit_and_liability_report(book: Workbook, year: int = REPORT_YEAR) -> None:
    """Create the annual management report consumed by FinanceRefresh.bas."""
    sheet = book.create_sheet(f'{year}年利润和负债')
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'C5'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = 'B2:AJ48'

    month_columns = {1: 'C', 2: 'E', 3: 'G', 4: 'K', 5: 'M', 6: 'O', 7: 'S', 8: 'U', 9: 'W', 10: 'AA', 11: 'AC', 12: 'AE'}
    percent_columns = {1: 'D', 2: 'F', 3: 'H', 4: 'L', 5: 'N', 6: 'P', 7: 'T', 8: 'V', 9: 'X', 10: 'AB', 11: 'AD', 12: 'AF'}
    quarter_columns = {1: ('I', (1, 2, 3)), 2: ('Q', (4, 5, 6)), 3: ('Y', (7, 8, 9)), 4: ('AG', (10, 11, 12))}
    quarter_percent_columns = {1: 'J', 2: 'R', 3: 'Z', 4: 'AH'}
    amount_rows = {
        5: '收入', 7: '成本', 9: '费用', 10: '   管理费用', 11: '其中：研发费用',
        12: '   财务费用', 13: '   销售费用', 15: '利润总额', 17: '所得税',
    }
    balance_rows = {23: '资产', 25: '应收账款', 27: '应付账款', 29: '负债',
                    31: '未分配利润', 33: '所有者权益'}
    revenue_rows = {45: '收入', 46: '项目收入', 47: '研发收入', 48: '销售商品收入'}

    widths = {'A': 3, 'B': 20, 'AJ': 11, 'AK': 2}
    for column in (*month_columns.values(), *(v[0] for v in quarter_columns.values()), 'AI'):
        widths[column] = 15
    for column in (*percent_columns.values(), *quarter_percent_columns.values()):
        widths[column] = 9
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions['AK'].hidden = True

    dark = PatternFill('solid', fgColor='263D56')
    blue = PatternFill('solid', fgColor='D9EAF7')
    light = PatternFill('solid', fgColor='EDF2F7')
    white_bold = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    section_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    thin_gray = Side(style='thin', color='C9D2DC')
    bottom_border = Border(bottom=thin_gray)
    for row in sheet.iter_rows(min_row=1, max_row=48, min_col=1, max_col=37):
        for cell in row:
            cell.font = Font(name='Arial', size=10)

    sheet['B2'] = '表1'
    sheet.merge_cells('B3:AI3')
    sheet['B3'] = f'一、{year}年利润表分析'
    sheet.merge_cells('AJ3:AJ4')
    sheet['AJ3'] = '趋势'
    sheet.merge_cells('B21:AI21')
    sheet['B21'] = f'二、{year}年资产负债表项目'
    sheet.merge_cells('AJ21:AJ22')
    sheet['AJ21'] = '趋势'
    sheet['B42'] = '表3'
    sheet.merge_cells('B43:AI43')
    sheet['B43'] = '各项收入占比'
    for address in ('B3:AJ3', 'B21:AJ21', 'B43:AI43'):
        for row in sheet[address]:
            for cell in row:
                cell.fill = dark
                cell.font = section_font
                cell.alignment = Alignment(vertical='center')

    headers = ['项目', '1月', '%', '2月', '%', '3月', '%', 'Q1季度', '%',
               '4月', '%', '5月', '%', '6月', '%', 'Q2季度', '%',
               '7月', '%', '8月', '%', '9月', '%', 'Q3季度', '%',
               '10月', '%', '11月', '%', '12月', '%', 'Q4季度', '%', 'YTD']
    for row_number in (4, 44):
        for offset, value in enumerate(headers, 2):
            cell = sheet.cell(row_number, offset, value)
            cell.fill = dark
            cell.font = white_bold
            cell.alignment = Alignment(horizontal='center', vertical='center')
    balance_headers = ['项目', '1月', '', '2月', '', '3月', '', 'Q1季度', '',
                       '4月', '', '5月', '', '6月', '', 'Q2季度', '',
                       '7月', '', '8月', '', '9月', '', 'Q3季度', '',
                       '10月', '', '11月', '', '12月', '', 'Q4季度']
    for offset, value in enumerate(balance_headers, 2):
        cell = sheet.cell(22, offset, value)
        cell.fill = dark
        cell.font = white_bold
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_number, label in amount_rows.items():
        sheet.cell(row_number, 2, label)
    for row_number, label in balance_rows.items():
        sheet.cell(row_number, 2, label)
    for row_number, label in revenue_rows.items():
        sheet.cell(row_number, 2, label)
    for row_number in (*amount_rows, *balance_rows, 35, 37, *revenue_rows):
        sheet.cell(row_number, 2).font = Font(name='Arial', size=10, bold=True)
        for column in range(2, 37):
            sheet.cell(row_number, column).border = bottom_border
    sheet['B35'] = '毛利率'
    sheet['B37'] = '资产负债率'

    for month, amount_column in month_columns.items():
        percent_column = percent_columns[month]
        # The macro writes source amounts into these month cells.
        for row_number in (5, 7, 10, 11, 12, 13, 15, 17, 23, 25, 27, 29, 31, 33, 46, 47, 48):
            sheet[f'{amount_column}{row_number}'].fill = blue
            sheet[f'{amount_column}{row_number}'].number_format = MONEY
        sheet[f'{amount_column}9'] = (
            f'=IF(COUNT({amount_column}10,{amount_column}12,{amount_column}13)=0,"",'
            f'SUM({amount_column}10,{amount_column}12,{amount_column}13))'
        )
        sheet[f'{amount_column}45'] = f'=IF({amount_column}5="","",{amount_column}5)'
        sheet[f'{amount_column}35'] = (
            f'=IF(OR({amount_column}5="",{amount_column}5=0,{amount_column}7=""),"",'
            f'({amount_column}5-{amount_column}7)/{amount_column}5)'
        )
        sheet[f'{amount_column}37'] = (
            f'=IF(OR({amount_column}23="",{amount_column}23=0,{amount_column}29=""),"",'
            f'{amount_column}29/{amount_column}23)'
        )
        for row_number in amount_rows:
            sheet[f'{percent_column}{row_number}'] = (
                f'=IF(OR({amount_column}{row_number}="",{amount_column}$5="",{amount_column}$5=0),"",'
                f'{amount_column}{row_number}/{amount_column}$5)'
            )
        for row_number in revenue_rows:
            sheet[f'{percent_column}{row_number}'] = (
                f'=IF(OR({amount_column}{row_number}="",{amount_column}$45="",{amount_column}$45=0),"",'
                f'{amount_column}{row_number}/{amount_column}$45)'
            )

    for quarter, (quarter_column, months) in quarter_columns.items():
        source_columns = [month_columns[month] for month in months]
        percent_column = quarter_percent_columns[quarter]
        for row_number in amount_rows:
            refs = ','.join(f'{column}{row_number}' for column in source_columns)
            sheet[f'{quarter_column}{row_number}'] = f'=IF(COUNT({refs})=0,"",SUM({refs}))'
            sheet[f'{percent_column}{row_number}'] = (
                f'=IF(OR({quarter_column}{row_number}="",{quarter_column}$5="",{quarter_column}$5=0),"",'
                f'{quarter_column}{row_number}/{quarter_column}$5)'
            )
        quarter_end = source_columns[-1]
        for row_number in balance_rows:
            sheet[f'{quarter_column}{row_number}'] = f'=IF({quarter_end}{row_number}="","",{quarter_end}{row_number})'
        sheet[f'{quarter_column}35'] = (
            f'=IF(OR({quarter_column}5="",{quarter_column}5=0,{quarter_column}7=""),"",'
            f'({quarter_column}5-{quarter_column}7)/{quarter_column}5)'
        )
        sheet[f'{quarter_column}37'] = (
            f'=IF(OR({quarter_column}23="",{quarter_column}23=0,{quarter_column}29=""),"",'
            f'{quarter_column}29/{quarter_column}23)'
        )
        for row_number in revenue_rows:
            refs = ','.join(f'{column}{row_number}' for column in source_columns)
            sheet[f'{quarter_column}{row_number}'] = f'=IF(COUNT({refs})=0,"",SUM({refs}))'
            sheet[f'{percent_column}{row_number}'] = (
                f'=IF(OR({quarter_column}{row_number}="",{quarter_column}$45="",{quarter_column}$45=0),"",'
                f'{quarter_column}{row_number}/{quarter_column}$45)'
            )

    for row_number in (*amount_rows, *revenue_rows):
        sheet[f'AI{row_number}'] = f'=IF(COUNT(I{row_number},Q{row_number},Y{row_number},AG{row_number})=0,"",SUM(I{row_number},Q{row_number},Y{row_number},AG{row_number}))'
    for row_number in (35, 37):
        for column in (*month_columns.values(), *(value[0] for value in quarter_columns.values())):
            sheet[f'{column}{row_number}'].number_format = '0.0%'
    for row_number in (*amount_rows, *revenue_rows):
        for column in (*percent_columns.values(), *quarter_percent_columns.values()):
            sheet[f'{column}{row_number}'].number_format = '0.0%'
    for row_number in (*amount_rows, *balance_rows, *revenue_rows):
        for column in (*month_columns.values(), *(v[0] for v in quarter_columns.values()), 'AI'):
            sheet[f'{column}{row_number}'].number_format = MONEY
    for row_number in range(5, 49):
        if row_number not in (21, 22, 43, 44):
            for column in range(2, 37):
                if sheet.cell(row_number, column).fill.fill_type is None:
                    sheet.cell(row_number, column).fill = light
    sheet['AK1'] = ''
    sheet['AK2'] = ''
    sheet['AK3'] = ''


def style_range(sheet, address, *, number_format=None, fill=None, color=None):
    for row in sheet[address]:
        for cell in row:
            if number_format:
                cell.number_format = number_format
            if fill:
                cell.fill = PatternFill('solid', fgColor=fill)
            if color:
                font = copy(cell.font)
                font.color = color
                cell.font = font


def build_workbook() -> Workbook:
    """Create from declarative labels and formula code; never read an existing XLSX."""
    layout = json.loads(files(__package__).joinpath('template_layout.json').read_text(encoding='utf-8'))
    book = Workbook()
    book.remove(book.active)
    book.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    for name, definition in layout.items():
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows(min_row=1, max_row=40, max_col=14):
            for cell in row:
                cell.font = Font(name='Arial', size=11)
        for row in range(1, 41):
            sheet.row_dimensions[row].height = 23
        for column, width in definition['widths'].items():
            sheet.column_dimensions[column].width = width
        for address, value in definition['cells'].items():
            sheet[address] = value
            sheet[address].data_type = 's' if isinstance(value, str) else sheet[address].data_type
        sheet['A1'].font = Font(name='Arial', size=18, bold=True)
        if name != '控制台':
            for cell in sheet[4]:
                if cell.value is not None:
                    cell.fill = PatternFill('solid', fgColor='263D56')
                    cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            sheet.row_dimensions[4].height = 30
            sheet.freeze_panes = 'A5'

    # Preserve raw-table identifiers and numeric column formats used by the refresh macro.
    text_columns = {'刷新信息': 'AB', '公司列表': 'AB', '利润表': 'ABF', '资产负债表': 'AD',
                    '现金流量表': 'ABC', '科目余额': 'AB', '凭证明细': 'ACDEFGHKM',
                    '出纳账': 'ACDEFGHKM', '往来余额': 'ABCD', '月度趋势': 'ABCE'}
    for name, columns in text_columns.items():
        sheet = book[name]
        last_col = max(c.column for c in sheet[4] if c.value is not None)
        for row in sheet.iter_rows(min_row=5, max_row=max(5, max(c.row for row in sheet for c in row if c.value is not None)), max_col=last_col):
            for cell in row:
                cell.number_format = '@' if cell.column_letter in columns else MONEY
        if name in ('凭证明细', '出纳账'):
            sheet['C5'].number_format = 'yyyy-mm-dd'

    home = book['控制台']
    style_range(home, 'B4:B5', number_format='@', fill='FFF2CC', color='0000FF')
    home['B8'].number_format = 'yyyy-mm-dd hh:mm:ss'
    home['B6'] = '=IF(OR(B4="",B5=""),"请填写公司编号和月份",IF(AND(B4=\'刷新信息\'!B6,B5=\'刷新信息\'!B8),"当前选择与数据一致","选择已变化，请刷新"))'
    stats = book['经营统计']
    metrics = [
        'SUMIFS(\'利润表\'!C5:C200,\'利润表\'!F5:F200,"营业收入")',
        'SUMIFS(\'利润表\'!C5:C200,\'利润表\'!F5:F200,"营业成本")',
        'SUMIFS(\'利润表\'!C5:C200,\'利润表\'!F5:F200,"净利润")',
        'SUMIFS(\'出纳账\'!I5:I6000,\'出纳账\'!M5:M6000,"<>内部划转")',
        'SUMIFS(\'出纳账\'!J5:J6000,\'出纳账\'!M5:M6000,"<>内部划转")', 'B8-B9',
        'SUMIFS(\'往来余额\'!I5:I6000,\'往来余额\'!A5:A6000,"应收")-SUMIFS(\'往来余额\'!J5:J6000,\'往来余额\'!A5:A6000,"应收")',
        'SUMIFS(\'往来余额\'!J5:J6000,\'往来余额\'!A5:A6000,"应付")-SUMIFS(\'往来余额\'!I5:I6000,\'往来余额\'!A5:A6000,"应付")',
    ]
    for row, expression in enumerate(metrics, 5):
        stats[f'B{row}'] = f'=IF({MATCHES},{expression},"")'
    for month in range(1, 13):
        row = month + 15
        stats[f'A{row}'] = f'=IF(\'控制台\'!B5="","",LEFT(\'控制台\'!B5,4)&"-{month:02}")'
        for column, label in zip('BCD', ['营业收入', '营业成本', '净利润']):
            stats[f'{column}{row}'] = f'=IF(AND({MATCHES},COUNTIF(\'月度趋势\'!A5:A1000,A{row})>0),SUMIFS(\'月度趋势\'!D5:D1000,\'月度趋势\'!A5:A1000,A{row},\'月度趋势\'!E5:E1000,"{label}"),"")'
    style_range(stats, 'B5:B12', number_format=MONEY)
    style_range(stats, 'B16:D27', number_format=MONEY)
    chart = LineChart()
    chart.title = '营业收入与成本（元）'
    chart.add_data(Reference(stats, min_col=2, max_col=3, min_row=15, max_row=27), titles_from_data=True)
    chart.set_categories(Reference(stats, min_col=1, min_row=16, max_row=27))
    chart.legend.position = 'b'
    chart.display_blanks = 'gap'
    chart.anchor = TwoCellAnchor(_from=AnchorMarker(col=5, row=3), to=AnchorMarker(col=14, row=19))
    stats.add_chart(chart)

    budget_input = book['预算输入']
    style_range(budget_input, 'A5:F204', fill='FFF8E5', color='0000FF')
    style_range(budget_input, 'A1:C504', number_format='@')
    style_range(budget_input, 'E5:E204', number_format=MONEY)
    budget = book['预算对比']
    for row in range(5, 37):
        r = row
        for col in 'AB':
            budget[f'{col}{r}'] = f'=IF(\'利润表\'!{col}{r}="","",\'利润表\'!{col}{r})'
        budget[f'C{r}'] = f'=IF(AND({MATCHES},ISNUMBER(\'利润表\'!C{r})),\'利润表\'!C{r},"")'
        budget[f'D{r}'] = f'=IF(AND({MATCHES},COUNTIFS(\'预算输入\'!A5:A204,\'控制台\'!B4,\'预算输入\'!B5:B204,\'控制台\'!B5,\'预算输入\'!C5:C204,A{r})>0),SUMIFS(\'预算输入\'!E5:E204,\'预算输入\'!A5:A204,\'控制台\'!B4,\'预算输入\'!B5:B204,\'控制台\'!B5,\'预算输入\'!C5:C204,A{r}),"")'
        budget[f'E{r}'] = f'=IF(OR(D{r}="",C{r}=""),"",C{r}-D{r})'
        budget[f'F{r}'] = f'=IF(OR(D{r}="",D{r}=0),"",E{r}/ABS(D{r}))'
    style_range(budget, 'C5:E36', number_format=MONEY)
    style_range(budget, 'F5:F36', number_format='0.0%')

    age = book['账龄输入']
    style_range(age, 'A5:H504', fill='FFF8E5', color='0000FF')
    style_range(age, 'A1:E504', number_format='@')
    style_range(age, 'F5:G504', number_format='yyyy-mm-dd')
    style_range(age, 'H5:H504', number_format=MONEY)
    validation = DataValidation(type='list', formula1='"应收,应付"', allow_blank=True)
    validation.showErrorMessage = True
    validation.showDropDown = False
    age.add_data_validation(validation)
    validation.add('C5:C504')
    for r in range(5, 505):
        age[f'I{r}'] = f'=IF(OR(A{r}<> \'控制台\'!B4,B{r}<> \'控制台\'!B5,F{r}=""),"",{AS_OF}-F{r})'
        age[f'J{r}'] = f'=IF(OR(A{r}<> \'控制台\'!B4,B{r}<> \'控制台\'!B5,G{r}=""),"",MAX(0,{AS_OF}-G{r}))'
        age[f'K{r}'] = f'=IF(A{r}="","",IF(AND({MATCHES},A{r}=\'控制台\'!B4,B{r}=\'控制台\'!B5,OR(C{r}="应收",C{r}="应付"),COUNTIFS(\'往来余额\'!A5:A6000,C{r},\'往来余额\'!B5:B6000,D{r})=1,E{r}<>"",COUNTIFS(A$5:A$504,A{r},B$5:B$504,B{r},C$5:C$504,C{r},D$5:D$504,D{r},E$5:E$504,E{r})=1,ISNUMBER(F{r}),I{r}>=0,ISNUMBER(H{r}),H{r}>=0),"纳入","待核实"))'
    aging = book['账龄统计']
    for r in (5, 6):
        aging[f'B{r}'] = f'=IF({MATCHES},\'经营统计\'!B{r+6},"")'
        base = f'\'账龄输入\'!H5:H504,\'账龄输入\'!C5:C504,A{r},\'账龄输入\'!K5:K504,"纳入"'
        aging[f'C{r}'] = f'=IF({MATCHES},SUMIFS({base}),"")'
        aging[f'D{r}'] = f'=IF({MATCHES},B{r}-C{r},"")'
        for col, (low, high) in zip('EFGHIJ', [(0, 30), (31, 60), (61, 90), (91, 180), (181, 365), (366, 999999)]):
            aging[f'{col}{r}'] = f'=IF({MATCHES},SUMIFS({base},\'账龄输入\'!I5:I504,">={low}",\'账龄输入\'!I5:I504,"<={high}"),"")'
        aging[f'K{r}'] = f'=IF({MATCHES},SUMIFS({base},\'账龄输入\'!J5:J504,">0"),"")'
    style_range(aging, 'B5:K6', number_format=MONEY)
    guide = book['使用说明']
    for r in range(5, 16):
        guide.row_dimensions[r].height = 42
        guide[f'B{r}'].alignment = Alignment(wrap_text=True)
    add_annual_analysis_data(book)
    add_profit_and_liability_report(book)
    add_dashboard(book, MATCHES)
    add_rd_report(book, MATCHES)
    add_hightech_report(book, MATCHES)
    return book


def save_template(output: Path, *, overwrite: bool = False) -> Path:
    """Finish the workbook in a sibling temporary file before publishing it."""
    output = output.resolve()
    if output.suffix.lower() != '.xlsx':
        raise ValueError('Output must end in .xlsx')
    if output.exists() and not overwrite:
        raise FileExistsError('Output already exists. Use --overwrite to replace it.')
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix='.finance-template-', suffix='.xlsx', dir=output.parent)
    os.close(fd)
    temporary = Path(name)
    book = None
    try:
        book = build_workbook()
        book.save(temporary)
        if overwrite:
            os.replace(temporary, output)
        else:
            # Exclusive creation prevents a concurrent builder from overwriting a workbook.
            try:
                with output.open('xb') as destination, temporary.open('rb') as source:
                    import shutil
                    try:
                        shutil.copyfileobj(source, destination)
                    except OSError:
                        destination.close()
                        output.unlink(missing_ok=True)
                        raise
            except FileExistsError:
                raise FileExistsError('Output already exists. Use --overwrite to replace it.') from None
        return output
    finally:
        if book is not None:
            book.close()
        temporary.unlink(missing_ok=True)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description='Build a blank finance XLSX template without login or Node.js.')
    parser.add_argument('--output', type=Path, default=None, help='Output path; defaults to WORKSPACE/excel/finance-template.xlsx')
    parser.add_argument('--overwrite', action='store_true', help='Replace the existing output workbook')
    args = parser.parse_args(argv)
    output = args.output or project_root() / 'excel/finance-template.xlsx'
    try:
        result = save_template(output, overwrite=args.overwrite)
    except (OSError, ValueError) as exc:
        print(f'Template generation failed: {exc}')
        return 2
    print(f'Created {result}')
    print('Excel recalculates formulas when opened. Use the Excel installer to add the refresh button.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
