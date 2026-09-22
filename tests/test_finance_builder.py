"""Portable template generation, refresh contract and output protection."""
import os
from pathlib import Path
import subprocess
import sys
from unittest.mock import patch

from openpyxl import load_workbook
import pytest

from kdzwy_receipt_uploader.finance.build_template import build_workbook, save_template
from kdzwy_receipt_uploader.finance.snapshot import MANAGED_SHEETS

ROOT = Path(__file__).resolve().parents[1]


def test_generated_workbook_contains_all_reports(tmp_path):
    path = save_template(tmp_path / 'finance.xlsx')
    actual = load_workbook(path)
    assert len(actual.sheetnames) == 22
    assert actual.sheetnames[-1] == '2026年利润和负债'
    assert actual['年度分析数据'].sheet_state == 'hidden'
    report = actual['2026年利润和负债']
    assert report['B3'].value == '一、2026年利润表分析'
    assert report['B21'].value == '二、2026年资产负债表项目'
    assert report['I5'].value == '=IF(COUNT(C5,E5,G5)=0,"",SUM(C5,E5,G5))'
    assert report['AI5'].value == '=IF(COUNT(I5,Q5,Y5,AG5)=0,"",SUM(I5,Q5,Y5,AG5))'
    assert report['AE4'].value == '12月'
    assert report['AG5'].value == '=IF(COUNT(AA5,AC5,AE5)=0,"",SUM(AA5,AC5,AE5))'
    assert report['D46'].value == '=IF(OR(C46="",C$45="",C$45=0),"",C46/C$45)'
    assert report['C5'].value is None and report['W48'].value is None
    actual.close()


def test_refresh_contract_and_editable_ranges():
    book = build_workbook()
    assert set(MANAGED_SHEETS) <= set(book.sheetnames)
    assert book['刷新信息']['B5'].value == '2'
    assert book['控制台']['B4'].value is None and book['控制台']['B5'].value is None
    assert book['预算输入']['E204'].fill.fgColor.rgb.endswith('FFF8E5')
    assert book['账龄输入']['H504'].fill.fgColor.rgb.endswith('FFF8E5')
    validation, = book['账龄输入'].data_validations.dataValidation
    assert str(validation.sqref) == 'C5:C504'
    assert validation.formula1 == '"应收,应付"'
    chart, = book['经营统计']._charts
    assert [s.val.numRef.f for s in chart.ser] == ["'经营统计'!$B$16:$B$27", "'经营统计'!$C$16:$C$27"]
    assert book.calculation.fullCalcOnLoad and book.calculation.calcMode == 'auto'
    assert all(book[name].freeze_panes == 'A5' for name in MANAGED_SHEETS)
    book.close()


def test_refuses_existing_output_and_failed_build_preserves_it(tmp_path):
    output = tmp_path / 'existing.xlsx'
    output.write_bytes(b'original user workbook')
    with pytest.raises(FileExistsError):
        save_template(output)
    with patch('kdzwy_receipt_uploader.finance.build_template.build_workbook', side_effect=RuntimeError('failed')):
        with pytest.raises(RuntimeError):
            save_template(output, overwrite=True)
    assert output.read_bytes() == b'original user workbook'
    assert not list(tmp_path.glob('.finance-template-*'))
    save_template(output, overwrite=True)
    book = load_workbook(output)
    assert len(book.sheetnames) == 22
    book.close()


def test_no_output_written_with_wrong_extension(tmp_path):
    with pytest.raises(ValueError, match='xlsx'):
        save_template(tmp_path / 'template.xlsm')
    assert not list(tmp_path.iterdir())


def test_cli_builds_from_empty_workspace_without_node_or_existing_template(tmp_path):
    (tmp_path / 'config').mkdir()
    env = {**os.environ, 'KDZWY_PROJECT_ROOT': str(tmp_path), 'PYTHONPATH': str(ROOT / 'src'), 'PYTHONUTF8': '1'}
    command = [sys.executable, '-m', 'kdzwy_receipt_uploader.finance.build_template']
    result = subprocess.run(command, cwd=tmp_path, env=env, capture_output=True, text=True, timeout=20)
    assert result.returncode == 0, result.stderr + result.stdout
    target = tmp_path / 'excel/finance-template.xlsx'
    original = target.read_bytes()
    retry = subprocess.run(command, cwd=tmp_path, env=env, capture_output=True, text=True, timeout=20)
    assert retry.returncode == 2
    assert target.read_bytes() == original
    assert not (tmp_path / 'runtime').exists()
    assert not (tmp_path / 'http_sessions').exists()
