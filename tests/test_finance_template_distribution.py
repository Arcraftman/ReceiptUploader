"""A clean checkout includes a usable, data-free finance workbook."""
from pathlib import Path

from openpyxl import load_workbook

from kdzwy_receipt_uploader.finance.snapshot import MANAGED_SHEETS

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'excel/finance-template.xlsx'


def test_distributed_template_has_structure_without_business_data():
    book = load_workbook(TEMPLATE, data_only=True)
    assert len(book.sheetnames) == 18
    assert set(MANAGED_SHEETS) <= set(book.sheetnames)
    assert book['控制台']['B4'].value is None
    assert book['控制台']['B5'].value is None
    assert book['控制台']['B7'].value == '尚未读取数据'
    assert book['刷新信息']['B5'].value == '1'
    for name in MANAGED_SHEETS:
        if name == '刷新信息':
            assert all(book[name].cell(row, 2).value is None for row in range(6, 10))
        else:
            assert all(cell.value is None for row in book[name].iter_rows(min_row=5) for cell in row)
    for name, last_column in [('预算输入', 6), ('账龄输入', 8)]:
        assert all(cell.value is None for row in book[name].iter_rows(min_row=5, max_col=last_column) for cell in row)
    assert not [(s.title, c.coordinate) for s in book for row in s for c in row if c.data_type == 'e']
    assert all(book['经营统计'].cell(row, 2).value is None for row in range(5, 13))
    assert book.sheetnames[-1] == '2026年利润和负债'
    book.close()


def test_template_preserves_formula_and_installer_contract():
    book = load_workbook(TEMPLATE)
    formulas = [c.value for s in book for row in s for c in row if c.data_type == 'f']
    assert len(formulas) > 1600
    assert 'B4<>""' in book['经营统计']['B5'].value
    assert book['经营统计']._charts
    book.close()
    installer = (ROOT / 'scripts/finance/install-excel.ps1').read_text(encoding='utf-8')
    assert "Join-Path $PSScriptRoot '../../excel/finance-template.xlsx'" in installer
    assert "Join-Path $PSScriptRoot '../../excel/finance.xlsm'" in installer
    assert 'SaveAs($outputPath, 52)' in installer
    macro = (ROOT / 'excel/FinanceRefresh.bas').read_text(encoding='utf-8')
    assert 'UpdateProfitAndLiabilityReport report, company, period' in macro
    assert 'FindTrendValue' in macro
    assert 'FindBalanceValue' in macro
    assert 'FindSubjectCredit' in macro
