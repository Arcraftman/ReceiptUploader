"""Dashboard persistence contract; arithmetic is checked in desktop Excel."""
from openpyxl import load_workbook
from kdzwy_receipt_uploader.finance.build_template import save_template
from kdzwy_receipt_uploader.finance.snapshot import MANAGED_SHEETS


def test_dashboard_round_trip_preserves_native_charts_and_scoped_ranges(tmp_path):
    book = load_workbook(save_template(tmp_path / 'dashboard.xlsx'))
    sheet = book['财务分析看板']
    assert book.sheetnames[1] == sheet.title
    assert sheet.title not in MANAGED_SHEETS
    assert len(sheet._charts) == 6
    assert len(sheet.defined_names) == 16
    for chart in sheet._charts:
        assert chart.x_axis.delete is False
        assert chart.y_axis.delete is False
        for series in chart.ser:
            name = series.val.numRef.f.split('!')[1]
            assert name in sheet.defined_names
            assert '$R$1' in sheet.defined_names[name].attr_text
    # Missing values must use NA for chart helpers rather than plotting zero.
    assert sheet['T83'].value == '=IF(ISNUMBER(B83),B83,NA())'
    assert all(sheet[f'AJ{r}'].data_type=='f' for r in range(2,8))
    assert 'B4=$AJ$1' in sheet['AJ2'].value
    assert sheet.column_dimensions['T'].hidden
    assert book['控制台']['B4'].value is None
    assert book['月度趋势']['A5'].value is None
    book.close()
