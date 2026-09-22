"""Extract the approved sheet's visual layout, without business data or pivots."""
import argparse
from copy import copy, deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook


def extract(source, output):
    reference = load_workbook(source)
    original = reference['2026研发费用']
    book = Workbook()
    sheet = book.active
    sheet.title = original.title
    for row in original:
        for cell in row:
            target = sheet.cell(cell.row, cell.column)
            # Keep labels only. Monetary inputs, formulas and chart caches are removed.
            if isinstance(cell.value, str) and cell.data_type != 'f':
                target.value = cell.value
            for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
                setattr(target, attribute, copy(getattr(cell, attribute)))
            target.number_format = cell.number_format
    for key, dimension in original.column_dimensions.items():
        sheet.column_dimensions[key] = copy(dimension)
        sheet.column_dimensions[key].parent = sheet
        sheet.column_dimensions[key]._style = None
    for key, dimension in original.row_dimensions.items():
        sheet.row_dimensions[key] = copy(dimension)
        sheet.row_dimensions[key].parent = sheet
        sheet.row_dimensions[key]._style = None
    for attribute in ('sheet_format', 'sheet_properties', 'page_margins', 'page_setup', 'print_options'):
        setattr(sheet, attribute, copy(getattr(original, attribute)))
    for merged in original.merged_cells:
        sheet.merge_cells(str(merged))
    for original_chart in original._charts:
        chart = deepcopy(original_chart)
        chart.pivotSource = None
        for series in chart.ser:
            if series.val and series.val.numRef:
                series.val.numRef.numCache = None
            if series.cat and series.cat.numRef:
                series.cat.numRef.numCache = None
            if series.cat and series.cat.strRef:
                series.cat.strRef.strCache = None
            if series.tx and series.tx.strRef:
                series.tx.strRef.strCache = None
        sheet.add_chart(chart)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    reference.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    extract(args.source, args.output)
