"""One authorized API check with synthetic aggregate data, never credentials."""
import json
from pathlib import Path
from openpyxl import load_workbook
from kdzwy_receipt_uploader.finance.interpretation import interpret

root = Path('outputs/finance-ai-monitor').resolve()
book = load_workbook('outputs/finance-dashboard-rd/dashboard-final.xlsx',data_only=True)
sheet = book['财务分析看板']
rows = [[sheet.cell(r,c).value if sheet.cell(r,c).value != '' else None for c in range(1,17)] for r in range(83,91)]
payload = {'company':'company_1','month':'2026-08','rows':rows}
try:
    result = interpret(root,payload)
except ValueError as exc:
    print(str(exc))
    raise SystemExit(1)
(root/'live-result.json').write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
print('Generated six interpretations:', result['model'], [len(t) for t in result['texts']])
book.close()
